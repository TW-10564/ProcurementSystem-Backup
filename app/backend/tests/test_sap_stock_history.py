"""ingestion/sap_stock_history.py — the parser for manually-placed, dated
raw SAP ZM13 "Data"-sheet stock exports (found 2026-07-23: 9 weekly files
in data/sap/6月, data/sap/7月, a different shape than SAP_STOCK_GLOB's
SBRCODE/あだ名/SYOCODE/SYONAME/QTYの合計 files). Covers: content-based file
discovery (ignores files with the wrong sheet/headers), per-item
aggregation across multiple 保管場所 rows, filename-derived as_of_date
(including the "don't guess a future year" rule), and de-duplication by
as_of_date when the same date exists in two files on disk."""
from datetime import date

import pytest

from onefcst.ingestion.sap_stock_history import (
    find_stock_history_files, parse_stock_history_file, parse_all_stock_history)

HEADER = ["品目", "在庫数量", "単位", "保管場所", "S", "伝票", "明細", "Plnt",
         "品目グループ", "品目タイプ", "旧品目コード", "ロット管理区分", "品目コード",
         "安全在庫", "単位", "品目St", "在庫ゼロフラグ", "ValCl", "原価区分",
         "移動平均原価", "通貨", "受注原価", "通貨", "標準原価", "通貨",
         "品目階層", "品目階層", "在庫合計金額", "通貨", "品目テキスト(50)"]


def _row(code, qty, storage, text):
    r = [None] * len(HEADER)
    r[0], r[1], r[3], r[29] = code, qty, storage, text
    return r


def _write_zm13_workbook(path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "placeholder"
    ws = wb.create_sheet("Data")
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    wb.save(path)


def _write_unrelated_workbook(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["not", "the", "right", "shape"])
    wb.save(path)


def test_find_stock_history_files_matches_only_the_right_shape(tmp_path):
    good = tmp_path / "0601.xlsx"
    _write_zm13_workbook(good, [_row("111111", 10, "1P10", "Item A")])
    bad = tmp_path / "SAP部材管理_詳細版260714.xlsx"
    _write_unrelated_workbook(bad)

    hits = find_stock_history_files(tmp_path)
    assert hits == [good]


def test_parse_aggregates_across_storage_locations(tmp_path):
    path = tmp_path / "0601.xlsx"
    _write_zm13_workbook(path, [
        _row("111111", 908, "1P10", "Item A"),
        _row("111111", 700, "1W60", "Item A"),
        _row("222222", 25, "1P10", "Item B"),
    ])

    result = parse_stock_history_file(path, today=date(2026, 7, 23))
    assert result["as_of_date"] == date(2026, 6, 1)
    by_code = {r["code"]: r["qty"] for r in result["rows"]}
    assert by_code == {"111111": 1608.0, "222222": 25.0}


def test_as_of_date_from_filename_never_lands_in_the_future(tmp_path):
    # "0721" relative to "today"=2026-01-10 would be July 2026, in the
    # future — must roll back to the prior year instead of guessing forward.
    path = tmp_path / "0721.xlsx"
    _write_zm13_workbook(path, [_row("111111", 1, "1P10", "Item A")])

    result = parse_stock_history_file(path, today=date(2026, 1, 10))
    assert result["as_of_date"] == date(2025, 7, 21)


def test_filename_without_a_parseable_date_is_skipped(tmp_path):
    path = tmp_path / "stock_export_final.xlsx"
    _write_zm13_workbook(path, [_row("111111", 1, "1P10", "Item A")])

    assert parse_stock_history_file(path, today=date(2026, 7, 23)) is None


def test_same_date_across_two_files_is_deduped_first_match_wins(tmp_path):
    (tmp_path / "a_dir").mkdir()
    (tmp_path / "b_dir").mkdir()
    f1 = tmp_path / "a_dir" / "0601.xlsx"
    f2 = tmp_path / "b_dir" / "0601.xlsx"
    _write_zm13_workbook(f1, [_row("111111", 10, "1P10", "Item A")])
    _write_zm13_workbook(f2, [_row("111111", 999, "1P10", "Item A")])

    results = parse_all_stock_history(tmp_path, today=date(2026, 7, 23))
    assert len(results) == 1
    assert results[0]["as_of_date"] == date(2026, 6, 1)


def test_multiple_distinct_dates_all_kept_oldest_first(tmp_path):
    _write_zm13_workbook(tmp_path / "0601.xlsx", [_row("111111", 10, "1P10", "Item A")])
    _write_zm13_workbook(tmp_path / "0608.xlsx", [_row("111111", 20, "1P10", "Item A")])

    results = parse_all_stock_history(tmp_path, today=date(2026, 7, 23))
    assert [r["as_of_date"] for r in results] == [date(2026, 6, 1), date(2026, 6, 8)]
