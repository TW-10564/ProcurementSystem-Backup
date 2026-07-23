"""ingestion/sap_zp21.py — the parser for ZP21 "組込" (build-incorporation)
weekly exports (found 2026-07-23, data/sap/組込/0525_0531組込.xlsx etc).
Replaced SAP部材管理_詳細版 (出荷/shipment) as the WeeklyActual source for
forecast_views.plan_vs_actual per procurement's own explanation: they
track demand by build-incorporation timing, not PC-shipment timing.
Covers: content-based file discovery (disambiguated from
sap_stock_history.py's ZM13 files, which share the same "Data" sheet
name but different headers), per-file aggregation of duplicate item
rows, and rejection of a file whose rows disagree on week bounds."""
from datetime import date, datetime

from onefcst.ingestion.sap_zp21 import find_zp21_files, parse_zp21_file, parse_all_zp21

HEADER = ["品目", "品目テキスト", "品目階層", "テキスト", "組込数量", "単",
         "移動平均", "通貨", "合計金額", "通貨", "開始日", "終了日"]


def _row(code, qty, text):
    r = [None] * len(HEADER)
    r[0], r[1], r[4] = code, text, qty
    r[10] = datetime(2026, 6, 1)
    r[11] = datetime(2026, 6, 7)
    return r


def _write_zp21_workbook(path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "placeholder"
    ws = wb.create_sheet("Data")
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    wb.save(path)


def _write_zm13_stock_workbook(path):
    """Same sheet NAME ("Data") as ZP21, but ZM13's own headers — must be
    correctly excluded by find_zp21_files, not confused for a ZP21 file."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "placeholder"
    ws = wb.create_sheet("Data")
    ws.append(["品目", "在庫数量", "単位", "保管場所", "品目テキスト(50)"])
    ws.append(["111111", 10, "PC", "1P10", "Item A"])
    wb.save(path)


def test_find_zp21_files_excludes_zm13_stock_files_with_the_same_sheet_name(tmp_path):
    good = tmp_path / "0601_0607組込.xlsx"
    _write_zp21_workbook(good, [_row("310637", 26, "CPU A")])
    stock = tmp_path / "0601.xlsx"
    _write_zm13_stock_workbook(stock)

    assert find_zp21_files(tmp_path) == [good]


def test_parse_aggregates_duplicate_item_rows_within_a_file(tmp_path):
    path = tmp_path / "0601_0607組込.xlsx"
    _write_zp21_workbook(path, [
        _row("310637", 26, "CPU A"),
        _row("310637", 4, "CPU A"),
        _row("281700", 129, "CPU B"),
    ])

    result = parse_zp21_file(path)
    assert result["week_start"] == date(2026, 6, 1)
    assert result["week_end"] == date(2026, 6, 7)
    by_code = {r["code"]: r["qty"] for r in result["rows"]}
    assert by_code == {"310637": 30.0, "281700": 129.0}


def test_file_with_disagreeing_week_bounds_across_rows_is_rejected(tmp_path):
    path = tmp_path / "bad.xlsx"
    rows = [_row("310637", 26, "CPU A")]
    rows[0][11] = datetime(2026, 6, 14)   # disagrees with the row's own start-week pairing
    _write_zp21_workbook(path, rows + [_row("281700", 1, "CPU B")])

    assert parse_zp21_file(path) is None


def test_multiple_weeks_all_kept_oldest_first(tmp_path):
    import openpyxl

    def write(path, start, end, code, qty):
        wb = openpyxl.Workbook()
        wb.active.title = "placeholder"
        ws = wb.create_sheet("Data")
        ws.append(HEADER)
        r = [None] * len(HEADER)
        r[0], r[4], r[10], r[11] = code, qty, start, end
        ws.append(r)
        wb.save(path)

    write(tmp_path / "a.xlsx", datetime(2026, 6, 1), datetime(2026, 6, 7), "310637", 26)
    write(tmp_path / "b.xlsx", datetime(2026, 6, 8), datetime(2026, 6, 14), "310637", 30)

    results = parse_all_zp21(tmp_path)
    assert [r["week_start"] for r in results] == [date(2026, 6, 1), date(2026, 6, 8)]
