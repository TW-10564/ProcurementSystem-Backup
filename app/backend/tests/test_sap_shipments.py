"""ingestion/sap_shipments.py — the parser for 【クエリ】部材利用(出荷), the
real per-item weekly SHIPPED quantity sheet that populates the Actual table
(found empty, 2026-07-15, before this module existed). Covers the parser
against a synthetic in-memory workbook mimicking the REAL file's exact
quirks (verified against data/sap/SAP部材管理_詳細版260707.xlsx before
writing this parser): 該当無し fallback, blank spacer rows, category 合計
subtotals, and the derived Windows/CPU-ratio tail block that must not be
parsed as real items. Also covers build._ingest_sap_shipments' resolution
+ Actual-write logic."""
from datetime import date, datetime

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from onefcst.models import Base, Item, ItemAlias, Actual, WeeklyActual, Snapshot
from onefcst.ingestion.sap_shipments import parse_sap_shipments, aggregate_by_month, flatten_weekly
from onefcst.registry.build import build_normalized_alias_index
from onefcst.build import _ingest_sap_shipments


@pytest.fixture()
def session():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine)()
    yield s
    s.close()


def _write_workbook(tmp_path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "製品名称読みかえ"   # required sheet name_crosswalk.py also reads
    ws2 = wb.create_sheet("【クエリ】部材利用(出荷)")
    for row in rows:
        ws2.append(row)
    path = tmp_path / "SAP部材管理_詳細版_test.xlsx"
    wb.save(path)
    return path


REAL_SHEET_SAMPLE = [
    (None, None, None, None, None, None),
    ("反映はB列から→", "Column2", "Column13", "Column14", "Column15", "Column16"),
    (None, None, None, None, None, None),
    (None, "raytrek", datetime(2026, 6, 2), datetime(2026, 6, 9),
     datetime(2026, 6, 16), datetime(2026, 6, 23)),
    ("↓読みかえ関数↓", "GALLERIA", datetime(2026, 6, 8), datetime(2026, 6, 15),
     datetime(2026, 6, 22), datetime(2026, 6, 29)),
    ("RTX5090", "5090", 13, 20, 20, 29),
    ("該当無し", "RTX3080 10G", 0, 0, 0, 0),   # 読みかえ found no match — use col B
    ("該当無し", "合計", 3044, 3269, 3194, 3484),   # category subtotal — must skip
    ("該当無し", None, None, None, None, None),    # blank separator — must skip
    ("U9 285K", "U9-285K", 33, 29, 36, 31),
    (None, "合計", 958, 1211, 1426, 1395),
    (None, "Windows11", 3418, 3542, 3994, 4516),   # derived tail — parsing must STOP here
    (None, "OEM", 3413, 3539, 3981, 4509),
    (None, "u9", 38, 35, 37, 36),
]


def test_parses_real_item_rows_and_skips_subtotals_and_blanks(tmp_path):
    path = _write_workbook(tmp_path, REAL_SHEET_SAMPLE)
    result = parse_sap_shipments(path)
    names = [(r["name_a"], r["name_b"]) for r in result["rows"]]
    assert ("RTX5090", "5090") in names
    assert ("該当無し", "RTX3080 10G") in names
    assert ("U9 285K", "U9-285K") in names
    # subtotal/blank/derived rows never make it into the parsed row list
    assert not any(nb == "合計" for _, nb in names)
    assert not any(na is None and nb is None for na, nb in names)
    assert not any(nb == "Windows11" for _, nb in names)


def test_stops_at_windows11_and_never_parses_the_derived_tail_block(tmp_path):
    path = _write_workbook(tmp_path, REAL_SHEET_SAMPLE)
    result = parse_sap_shipments(path)
    names_b = {r["name_b"] for r in result["rows"]}
    assert "OEM" not in names_b     # would be a dangerously generic label to resolve
    assert "u9" not in names_b      # derived CPU-tier breakdown, re-counts rows above


def test_week_ending_dates_and_weekly_quantities_are_read_correctly(tmp_path):
    path = _write_workbook(tmp_path, REAL_SHEET_SAMPLE)
    result = parse_sap_shipments(path)
    assert result["week_ending_dates"] == [date(2026, 6, 2), date(2026, 6, 9),
                                           date(2026, 6, 16), date(2026, 6, 23)]
    rtx5090 = next(r for r in result["rows"] if r["name_b"] == "5090")
    assert rtx5090["weekly_qty"] == [13.0, 20.0, 20.0, 29.0]


def test_aggregate_by_month_sums_weeks_landing_in_the_same_month(tmp_path):
    path = _write_workbook(tmp_path, REAL_SHEET_SAMPLE)
    result = parse_sap_shipments(path)
    agg = aggregate_by_month(result)
    rtx5090 = next(r for r in agg if r["name_b"] == "5090")
    assert rtx5090["months"] == {date(2026, 6, 1): 13.0 + 20.0 + 20.0 + 29.0}


def test_returns_none_when_the_sheet_is_missing(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "製品名称読みかえ"
    path = tmp_path / "no_shipment_sheet.xlsx"
    wb.save(path)
    assert parse_sap_shipments(path) is None


def test_returns_none_when_no_file_found(monkeypatch):
    import onefcst.ingestion.sap_shipments as mod
    monkeypatch.setattr(mod, "latest_crosswalk_file", lambda: None)
    assert parse_sap_shipments() is None


# ------------------------------------------------- build._ingest_sap_shipments

def test_ingest_writes_actual_rows_for_resolved_items_via_name_a(session):
    item = Item(canonical_name="RTX5090", kind="part")
    session.add(item)
    session.flush()
    session.add(ItemAlias(item_id=item.id, alias="RTX5090", source_system="ONE綽名"))
    session.commit()
    norm_index = build_normalized_alias_index(session)

    agg_rows = [{"name_a": "RTX5090", "name_b": "5090",
                "months": {date(2026, 6, 1): 82.0}}]
    stats = _ingest_sap_shipments(session, norm_index, agg_rows)
    session.commit()

    assert stats == {"resolved": 1, "unresolved": 0}
    row = session.query(Actual).one()
    assert row.item_id == item.id
    assert row.qty == 82.0
    assert row.source == "SAP"


def test_ingest_falls_back_to_name_b_when_name_a_is_該当無し(session):
    item = Item(canonical_name="RTX3080 10G", kind="part")
    session.add(item)
    session.flush()
    session.add(ItemAlias(item_id=item.id, alias="RTX3080 10G", source_system="ONE綽名"))
    session.commit()
    norm_index = build_normalized_alias_index(session)

    agg_rows = [{"name_a": "該当無し", "name_b": "RTX3080 10G",
                "months": {date(2026, 6, 1): 5.0}}]
    stats = _ingest_sap_shipments(session, norm_index, agg_rows)
    session.commit()

    assert stats == {"resolved": 1, "unresolved": 0}
    assert session.query(Actual).one().item_id == item.id


def test_ingest_skips_unresolvable_rows_without_writing_a_row(session):
    norm_index = build_normalized_alias_index(session)
    agg_rows = [{"name_a": "該当無し", "name_b": "Totally Unknown Part",
                "months": {date(2026, 6, 1): 5.0}}]
    stats = _ingest_sap_shipments(session, norm_index, agg_rows)
    session.commit()

    assert stats == {"resolved": 0, "unresolved": 1}
    assert session.query(Actual).count() == 0


# --------------------------------------------------------------- flatten_weekly

def test_flatten_weekly_keeps_raw_per_week_values_not_aggregated(tmp_path):
    path = _write_workbook(tmp_path, REAL_SHEET_SAMPLE)
    result = parse_sap_shipments(path)
    flat = flatten_weekly(result)
    rtx5090 = next(r for r in flat if r["name_b"] == "5090")
    assert rtx5090["weeks"] == {
        date(2026, 6, 2): 13.0, date(2026, 6, 9): 20.0,
        date(2026, 6, 16): 20.0, date(2026, 6, 23): 29.0,
    }


def test_flatten_weekly_and_aggregate_by_month_stay_row_aligned(tmp_path):
    """build._ingest_sap_shipments zips agg_rows with flatten_weekly's
    output by position — both must iterate parsed['rows'] in the same
    order for that alignment to be valid."""
    path = _write_workbook(tmp_path, REAL_SHEET_SAMPLE)
    result = parse_sap_shipments(path)
    agg = aggregate_by_month(result)
    flat = flatten_weekly(result)
    assert [r["name_b"] for r in agg] == [r["name_b"] for r in flat]


# ------------------------------------------------------- WeeklyActual ingestion

def test_ingest_writes_weekly_actual_rows_alongside_monthly_actual(session):
    item = Item(canonical_name="RTX5090", kind="part")
    session.add(item)
    session.flush()
    session.add(ItemAlias(item_id=item.id, alias="RTX5090", source_system="ONE綽名"))
    session.commit()
    norm_index = build_normalized_alias_index(session)

    agg_rows = [{"name_a": "RTX5090", "name_b": "5090",
                "months": {date(2026, 6, 1): 82.0}}]
    weekly_rows = [{"name_a": "RTX5090", "name_b": "5090",
                   "weeks": {date(2026, 6, 2): 13.0, date(2026, 6, 9): 20.0,
                            date(2026, 6, 16): 20.0, date(2026, 6, 23): 29.0}}]
    stats = _ingest_sap_shipments(session, norm_index, agg_rows, weekly_rows)
    session.commit()

    assert stats == {"resolved": 1, "unresolved": 0}
    assert session.query(Actual).one().qty == 82.0
    weekly = {r.week_start: r.qty for r in session.query(WeeklyActual).all()}
    assert weekly == {date(2026, 6, 2): 13.0, date(2026, 6, 9): 20.0,
                      date(2026, 6, 16): 20.0, date(2026, 6, 23): 29.0}
    assert all(r.item_id == item.id and r.source == "SAP"
              for r in session.query(WeeklyActual).all())


def test_ingest_without_weekly_rows_still_writes_monthly_actual_only(session):
    """weekly_rows is optional (build.py's own caller always passes it when
    a real sheet was parsed, but the parameter defaults to None so a
    caller — or an older test — that only cares about monthly Actual isn't
    forced to supply it)."""
    item = Item(canonical_name="RTX5090", kind="part")
    session.add(item)
    session.flush()
    session.add(ItemAlias(item_id=item.id, alias="RTX5090", source_system="ONE綽名"))
    session.commit()
    norm_index = build_normalized_alias_index(session)

    agg_rows = [{"name_a": "RTX5090", "name_b": "5090",
                "months": {date(2026, 6, 1): 82.0}}]
    _ingest_sap_shipments(session, norm_index, agg_rows)
    session.commit()

    assert session.query(Actual).count() == 1
    assert session.query(WeeklyActual).count() == 0


def test_two_source_rows_resolving_to_the_same_item_are_summed_not_duplicate_inserted(session):
    """Real bug found live 2026-07-15 against the actual sheet: 'RX7600XT'
    (matched via name_a) and a separate '該当無し'/'RX7600XT' row (matched
    via name_b fallback) both resolved to the same catalog item — genuine
    data, not a parsing error (the sheet tracks some parts on more than
    one line). Inserting one Actual/WeeklyActual row per SOURCE row
    violated the (item_id, month, source) unique constraint and crashed
    the entire rebuild. Must sum instead."""
    item = Item(canonical_name="RX7600XT", kind="part")
    session.add(item)
    session.flush()
    session.add(ItemAlias(item_id=item.id, alias="RX7600XT", source_system="ONE綽名"))
    session.commit()
    norm_index = build_normalized_alias_index(session)

    agg_rows = [
        {"name_a": "RX7600XT", "name_b": "7600XT",
         "months": {date(2026, 6, 1): 5.0}},
        {"name_a": "該当無し", "name_b": "RX7600XT",
         "months": {date(2026, 6, 1): 3.0}},
    ]
    weekly_rows = [
        {"name_a": "RX7600XT", "name_b": "7600XT",
         "weeks": {date(2026, 6, 1): 2.0, date(2026, 6, 8): 3.0}},
        {"name_a": "該当無し", "name_b": "RX7600XT",
         "weeks": {date(2026, 6, 1): 1.0, date(2026, 6, 8): 2.0}},
    ]

    stats = _ingest_sap_shipments(session, norm_index, agg_rows, weekly_rows)
    session.commit()   # must not raise IntegrityError

    assert stats == {"resolved": 2, "unresolved": 0}
    actual = session.query(Actual).one()
    assert actual.item_id == item.id
    assert actual.qty == 8.0   # 5.0 + 3.0, summed not duplicated

    weekly = {r.week_start: r.qty for r in session.query(WeeklyActual).all()}
    assert weekly == {date(2026, 6, 1): 3.0, date(2026, 6, 8): 5.0}
