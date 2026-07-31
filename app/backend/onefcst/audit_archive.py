"""Monthly audit-log report: export the finished month, then clear it.

    cd app/backend
    python3 -m onefcst.audit_archive --month 2026-07             # dry run
    python3 -m onefcst.audit_archive --month 2026-07 --commit    # export + purge

WHY THIS EXISTS
    The change history is a single unbounded list. It is the record of who
    changed what, so it must be kept — but on screen it becomes unreadable
    long before it becomes large, and the team files a monthly report anyway.
    So: one Excel per calendar month, then the month's rows leave the database
    and the feed starts clean.

THE DANGEROUS PART, AND WHAT GUARDS IT
    This deletes audit records. Everything below exists because of that:

      * **Dry run by default.** --commit is required to delete anything.
      * **Export first, verify, then delete.** The written workbook is re-opened
        and its data rows counted against the number of rows selected. A file
        that cannot be reopened, or is short, aborts the purge — a half-written
        file next to a wiped table is the one outcome worth engineering against.
      * **The current month is never purged** without --force. "Restart next
        month" means archiving a month that has finished; purging a live month
        would delete this morning's entries, which people still need on screen.
      * **A marker row survives the purge**, naming the file and the count. A
        gap in an audit trail must be explained by the trail itself, or the
        next person to read it cannot tell an archive from a deletion.

    After a purge the Excel is the ONLY copy of those entries (bar a pg_dump
    still inside its retention window). It has to be kept somewhere backed up
    — that is a filing responsibility, not something this code can enforce, so
    the report says so out loud every time it runs.
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

from .db import get_session
from .export.excel import export_audit_month
from .models import AuditLog


def month_bounds(year: int, month: int) -> tuple[datetime, datetime]:
    start = datetime(year, month, 1)
    end = datetime(year + (month == 12), (month % 12) + 1, 1)
    return start, end


def parse_month(text: str) -> tuple[int, int]:
    """Accepts 2026-07 / 2026/07 / 202607 — whichever the operator types."""
    cleaned = text.replace("/", "-").strip()
    if "-" in cleaned:
        y, m = cleaned.split("-", 1)
    else:
        y, m = cleaned[:4], cleaned[4:]
    year, mon = int(y), int(m)
    if not 1 <= mon <= 12:
        raise ValueError(f"month out of range: {text}")
    return year, mon


def _verify(path: Path, expected: int) -> tuple[bool, str]:
    """Re-open the workbook and count its data rows.

    Trusting the write to have succeeded is exactly the assumption that makes
    a purge unrecoverable. openpyxl is used here in read-only mode rather than
    reusing the in-memory workbook, because the question is whether the FILE
    on disk is readable, not whether the object was built correctly.
    """
    from openpyxl import load_workbook

    if not path.exists():
        return False, "the export file was not created"
    try:
        ws = load_workbook(path, read_only=True).active
        # Header block is 4 rows: title, period+count, blank, column headers.
        found = max(0, ws.max_row - 4)
    except Exception as exc:                                   # noqa: BLE001
        return False, f"the export file could not be reopened: {exc}"
    if found != expected:
        return False, f"the export holds {found} rows but {expected} were selected"
    return True, f"verified {found} rows in {path.name}"


def archive(year: int, month: int, commit: bool = False, force: bool = False,
            out_dir: Path | None = None, by: str = "audit-archive") -> dict:
    start, end = month_bounds(year, month)
    now = datetime.now()
    is_current = start <= now < end

    session = get_session()
    try:
        total = (session.query(AuditLog)
                 .filter(AuditLog.at >= start, AuditLog.at < end).count())
    finally:
        session.close()

    report: dict = {"month": f"{year}-{month:02d}", "rows": total,
                    "exported": None, "purged": 0, "status": "ok"}

    if total == 0:
        report["status"] = "empty"
        return report

    if is_current and not force:
        # Not an error: exporting a live month for a mid-month review is a
        # perfectly reasonable thing to do. Only the deletion is refused.
        path, count = export_audit_month(year, month, out_dir)
        report.update(exported=str(path), rows=count, status="exported-only",
                      note="current month — exported but NOT purged (use --force "
                           "only if you really mean to delete this month's live "
                           "entries)")
        return report

    path, count = export_audit_month(year, month, out_dir)
    report["exported"] = str(path)
    report["rows"] = count

    ok, detail = _verify(path, count)
    report["verification"] = detail
    if not ok:
        report["status"] = "verification-failed"
        return report                       # nothing deleted, by design

    if not commit:
        report["status"] = "dry-run"
        return report

    session = get_session()
    try:
        purged = (session.query(AuditLog)
                  .filter(AuditLog.at >= start, AuditLog.at < end).delete())
        # Written AFTER the delete so it is not itself caught by it, and dated
        # now rather than backdated into the archived month.
        session.add(AuditLog(
            who=by,
            what=f"{year}年{month:02d}月の変更履歴 {purged}件をアーカイブしました"
                 f"（{Path(path).name}）",
            after=str(path)))
        session.commit()
    finally:
        session.close()

    report["purged"] = purged
    return report


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--month", required=True,
                    help="the month to archive, e.g. 2026-07")
    ap.add_argument("--commit", action="store_true",
                    help="actually delete the month's rows after a verified export")
    ap.add_argument("--force", action="store_true",
                    help="allow purging the CURRENT (unfinished) month")
    ap.add_argument("--out-dir", type=Path, default=None)
    ap.add_argument("--by", default="audit-archive")
    args = ap.parse_args()

    year, month = parse_month(args.month)
    r = archive(year, month, commit=args.commit, force=args.force,
                out_dir=args.out_dir, by=args.by)

    print(f"month     : {r['month']}")
    print(f"rows      : {r['rows']}")
    print(f"exported  : {r['exported'] or '—'}")
    if "verification" in r:
        print(f"verified  : {r['verification']}")
    print(f"purged    : {r['purged']}")
    print(f"status    : {r['status']}")
    if r.get("note"):
        print(f"note      : {r['note']}")

    if r["status"] == "verification-failed":
        print("\nNOTHING WAS DELETED — the export could not be verified.")
        return 1
    if r["status"] == "dry-run":
        print("\nDRY RUN — the file was written but no rows were deleted.")
        print("Re-run with --commit to clear the month.")
    if r["purged"]:
        print("\nThe exported file is now the only copy of those entries.")
        print("Store it where it is backed up — a single copy on this host is not a backup.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
