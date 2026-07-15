"""Reads 購買計画 (PO) rows across all 17 ONE sub-files, auto-detecting
columns from the header row (数量/$価格/¥価格/納品日/IO/発注日/綽名/仕入先/
発注番号) so it's robust to per-file column shifts.

Copied in-house from the frozen consolidation/ POC's purchasing.py
2026-07-14 (verify/order_suggestions.py used to import it live via a
sys.path hack) — see ingestion/subfile_config.py's docstring for why the
active app must not depend on consolidation/ at runtime. Used only for
verify/order_suggestions.py's plausibility check against real recent POs;
not part of the write/ingest pipeline itself."""
from __future__ import annotations

from .subfile_config import FILES, BASE
from .excel import to_date, to_float, cell_str


def _detect_cols(header: dict) -> dict:
    """Map field -> column index from the 購買計画 header row."""
    col = {}
    for c, raw in header.items():
        if raw is None:
            continue
        h = str(raw).strip()
        if "数量" in h:
            col["qty"] = c
        elif "$" in h and ("価格" in h or "単価" in h):
            col["usd"] = c
        elif any(y in h for y in ("¥", "\\", "￥")) and ("価格" in h or "単価" in h):
            col["jpy"] = c
        elif "納品" in h and "済" not in h:
            col["delivery"] = c
        elif h == "IO":
            col["io"] = c
        elif "発注日" in h:
            col["order_date"] = c
        elif "綽名" in h:
            col["sku"] = c
        elif ("仕入先" in h or "購買先" in h) and "コード" not in h:
            col["supplier"] = c
        elif "発注書番号" in h or "発注番号" in h:
            col["po_num"] = c
    return col


def read_kobai(cat: str) -> list[dict]:
    """Return list of PO dicts for one category's 購買計画 (auto-detected layout)."""
    cfg = FILES[cat]
    path = BASE / cfg["file"]
    rows = []
    if cfg["fmt"] == "xlsb":
        import pyxlsb
        with pyxlsb.open_workbook(str(path)) as wb:
            if "購買計画" not in wb.sheets:
                return []
            with wb.get_sheet("購買計画") as ws:
                for r in ws.rows():
                    rows.append({c.c: c.v for c in r})
    else:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            if "購買計画" not in wb.sheetnames:
                return []
            ws = wb["購買計画"]
            for r in ws.iter_rows(values_only=True):
                rows.append({j: v for j, v in enumerate(r)})
        finally:
            wb.close()
    if not rows:
        return []

    col = _detect_cols(rows[0])
    if "sku" not in col or "delivery" not in col:
        return []   # can't use this file's PO sheet reliably

    out = []
    for row in rows[1:]:
        sku = cell_str(row.get(col["sku"]))
        if not sku:
            continue
        rowtext = " ".join(str(v) for v in row.values() if isinstance(v, str))
        delivered = "納品済" in rowtext
        cancelled = "キャンセル" in rowtext
        out.append({
            "category": cat, "file": cfg["file"], "sku": sku,
            "qty": to_float(row.get(col.get("qty"))),
            "usd": to_float(row.get(col.get("usd"))),
            "jpy": to_float(row.get(col.get("jpy"))),
            "delivery": to_date(row.get(col.get("delivery"))),
            "order_date": to_date(row.get(col.get("order_date"))),
            "io": to_float(row.get(col.get("io"))),
            "supplier": cell_str(row.get(col.get("supplier"))),
            "po_num": cell_str(row.get(col.get("po_num"))),
            "status": "delivered" if delivered else "cancelled" if cancelled else "open",
        })
    return out


def read_all_pos() -> list[dict]:
    pos = []
    for cat in FILES:
        try:
            pos += read_kobai(cat)
        except Exception:
            pass
    return pos
