"""Shared read-only Excel access helpers (xlsx/xlsm via openpyxl, xlsb via pyxlsb).

Guarantees:
- Sources are opened read-only, values-only. We never save a source workbook.
- Every reader returns rows plus enough position info to record lineage.

Missing-file/-sheet convention for ingestion/*.py parsers (standardized
2026-07-14 — previously each parser picked its own): a source the pipeline
CANNOT proceed without (①係数計算, ONE計画.xlsm, the 17 sub-files) should
let a missing file/sheet raise and fail loudly — see config.planning_file,
which already does this. A source that's genuinely OPTIONAL (SAP名称
crosswalk, 在庫の持ち方 policy table — a human may not always have supplied
one) should degrade gracefully to an empty/None result instead, per
read_xlsx_rows' required=False and ingestion/stock_policy.py's glob-miss
handling. Naming: parsers are named parse_*(), not load_*()/get_*(), except
where a caller needs to inject the source path directly for testability
(ingestion/mb_procurement.py's parse_mb_takeorpay(one_dir) — verify/
mismatch.py's tests monkeypatch this exact call to avoid touching the real
NAS file)."""
from __future__ import annotations

import hashlib
from datetime import datetime, timedelta
from pathlib import Path


def file_meta(path: Path) -> dict:
    st = path.stat()
    h = hashlib.md5()
    with open(path, "rb") as f:          # noqa: PTH123
        while chunk := f.read(1 << 20):
            h.update(chunk)
    return {"mtime": datetime.fromtimestamp(st.st_mtime), "checksum": h.hexdigest()}


def _calamine_sheet(path: Path, sheet: str | None, required: bool):
    """(rows, sheet_names) read via python-calamine, normalized to openpyxl's
    value conventions — or None if calamine can't serve this file.

    Why calamine at all: openpyxl materialises the whole style table before
    it will hand over a single value, and the ONE workbooks carry ~1M named
    styles, so `apply_stylesheet` alone was ~61% of a full build (profiled
    2026-07-28). calamine is a Rust reader that ignores formatting entirely
    — measured 8.5x faster on ONE_08_VGA's 部品詳細.

    Normalizations applied so callers see EXACTLY what openpyxl gave them:
      empty cell  calamine "" -> None   (openpyxl's convention; read_dict_rows
                  and every `if v is None` check downstream depend on it)
      date cell   calamine datetime.date -> datetime.datetime at midnight

    NOT equivalent in one respect: Excel ERROR cells (#N/A, #DIV/0!) come
    back as "" from calamine, i.e. None here, where openpyxl yields the
    literal string "#N/A". Those are never numeric plan values and are
    already discarded downstream — verified by a full-build diff showing
    byte-identical Gate A / registry / per-subfile cell counts either way.
    If a future caller ever needs to SEE error cells, it must use
    data_only=False (openpyxl) instead."""
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        return None
    from datetime import date as _date
    try:
        wb = CalamineWorkbook.from_path(str(path))
        names = list(wb.sheet_names)
        if sheet is None:
            sheet = names[0]
        elif sheet not in names:
            if required:
                raise KeyError(sheet)
            # rows=None distinguishes "sheet absent" from "sheet present but
            # empty" ([]), which read_xlsx_rows must report differently.
            return None, names
        # skip_empty_area=False is REQUIRED: the default trims leading empty
        # rows/columns, which would shift every absolute column index the
        # per-file layouts in subfile_config.py are pinned to.
        raw = wb.get_sheet_by_name(sheet).to_python(skip_empty_area=False)
    except KeyError:
        raise
    except Exception:
        return None                      # any reader problem -> openpyxl path
    rows = [tuple(_cal_value(v, _date) for v in r) for r in raw]
    return rows, names


def _cal_value(v, _date):
    """One calamine cell -> the value openpyxl would have produced.

    The int case is NOT cosmetic. calamine returns every number as f64, so a
    numeric 綽名 — VGA's SKUs are literally 5050/5060/5070/5080/5090 — comes
    back as 5090.0 and stringifies to "5090.0", which matches NOTHING in the
    alias registry. 綽名 is the join key for the entire system (see CLAUDE.md),
    so this silently orphans every numerically-named SKU. openpyxl yields int
    for whole numbers; mirror that."""
    if v == "":
        return None
    if type(v) is _date:
        return datetime(v.year, v.month, v.day)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def read_xlsx_rows(path: Path, sheet: str | None = None, max_row: int | None = None,
                   max_col: int | None = None, data_only: bool = True,
                   required: bool = True) -> list[tuple] | None:
    """data_only=True (default) returns cached VALUES (what Excel last computed).
    data_only=False returns FORMULA STRINGS instead, for cells that hold formulas
    (used by verify/parity.py's chain check to classify each cell's formula
    kind — see engine/chain.py's classify_formula).

    sheet=None reads the FIRST sheet (for callers like ingestion/sap.py whose
    source's sheet name isn't fixed/known ahead of time).

    required=False returns None instead of raising KeyError when `sheet` is
    given but absent — for genuinely optional sources (ingestion/
    name_crosswalk.py's crosswalk file may or may not have its expected
    sheet). Callers of genuinely-mandatory sources should leave this True so
    a missing sheet still fails loudly."""
    # data_only=False means the caller wants FORMULA STRINGS (plan1_chain's
    # chain check) — calamine only ever returns computed values, so that
    # request must stay on openpyxl.
    if data_only:
        got = _calamine_sheet(path, sheet, required)
        if got is not None:
            rows, _names = got
            if rows is None:              # sheet absent, required=False
                return None
            if max_row is not None:
                rows = rows[:max_row]
            if max_col is not None:
                rows = [r[:max_col] for r in rows]
            return rows

    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=data_only,
                                keep_vba=path.suffix == ".xlsm")
    try:
        if sheet is None:
            sheet = wb.sheetnames[0]
        elif sheet not in wb.sheetnames and not required:
            return None
        ws = wb[sheet]
        return [row for row in ws.iter_rows(max_row=max_row, max_col=max_col,
                                            values_only=True)]
    finally:
        wb.close()


def _xlsb_datetime(v):
    """pyxlsb returns Excel serials as floats; convert plausible dates."""
    if isinstance(v, float) and 20000 < v < 60000:
        return datetime(1899, 12, 30) + timedelta(days=v)
    return v


def read_xlsb_rows(path: Path, sheet: str, max_row: int | None = None,
                   max_col: int | None = None, dates: bool = False) -> list[tuple]:
    from pyxlsb import open_workbook
    out = []
    with open_workbook(str(path)) as wb:
        with wb.get_sheet(sheet) as ws:
            for ri, row in enumerate(ws.rows(), 1):
                if max_row and ri > max_row:
                    break
                vals = [c.v for c in (row[:max_col] if max_col else row)]
                if dates:
                    vals = [_xlsb_datetime(v) for v in vals]
                out.append(tuple(vals))
    return out


def to_date(v):
    """Excel serial / datetime / date -> date, or None. Shared value coercion
    (moved in-house from the frozen consolidation/ POC 2026-07-14 — see
    read_dict_rows's docstring for why)."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    from datetime import date as _date
    if isinstance(v, _date):
        return v
    try:
        f = float(v)
        if f <= 0:
            return None
        d = (datetime(1899, 12, 30) + timedelta(days=f)).date()
        return d if 2010 <= d.year <= 2035 else None
    except (TypeError, ValueError):
        return None


def to_float(v):
    """Safe float coercion — None/blank/error-string ('#REF!' etc) -> None."""
    if v is None:
        return None
    if isinstance(v, str) and (v.startswith("#") or v.strip() == ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def cell_str(v) -> str | None:
    """Trimmed string, or None for empty/None (avoids the str(None)=='None' trap)."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def read_dict_rows(path: Path, fmt: str, sheet: str) -> tuple[list[dict], bool, list[str]]:
    """Sparse dict-per-row reader ({col_index: value}, skipping None cells) —
    the shape ingestion/subfiles.py's block-parsing loop needs (row.get(col)
    lookups by absolute column index), as opposed to read_xlsx_rows'/
    read_xlsb_rows' positional-tuple shape.

    Returns (rows, sheet_present, all_sheet_names) so a caller can tell
    "sheet genuinely absent" apart from "sheet present but empty". Moved
    in-house from the frozen consolidation/ POC's consolidate.read_sheet
    2026-07-14 — the active app must not depend on the frozen one at
    runtime, only reuse its already-proven per-file layouts as a one-time
    reference (see ingestion/subfile_config.py)."""
    if fmt == "xlsb":
        import pyxlsb
        rows = []
        with pyxlsb.open_workbook(str(path)) as wb:
            names = wb.sheets
            if sheet not in names:
                return [], False, names
            with wb.get_sheet(sheet) as ws:
                for row in ws.rows():
                    rows.append({c.c: c.v for c in row if c.v is not None})
        return rows, True, names
    got = _calamine_sheet(path, sheet, required=False)
    if got is not None:
        rows_t, names = got
        if rows_t is None:
            return [], False, names
        rows = [{j: v for j, v in enumerate(row) if v is not None}
                for row in rows_t]
        return rows, True, names

    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        if sheet not in names:
            return [], False, names
        ws = wb[sheet]
        rows = [{j: v for j, v in enumerate(row) if v is not None}
               for row in ws.iter_rows(values_only=True)]
        return rows, True, names
    finally:
        wb.close()


def col_letter(idx0: int) -> str:
    """0-based column index -> Excel letters (for source_cell lineage)."""
    s, n = "", idx0 + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s
