"""SAP ZM13 weekly stock parser → physical stock per あだ名 (actuals source='SAP').

Layout (verified Phase-0): SBRCODE | SBRNAME | あだ名 | SYOCODE | SYONAME | QTYの合計.
Multiple SYOCODE rows per あだ名 are summed.

Blank-あだ名 rows (368/~900 in the file checked 2026-07-07, mostly MB) are
otherwise dropped entirely — meaning those items get no SAP stock linkage
and therefore no order suggestion, regardless of forecast quality. Resolved
via the procurement team's OWN maintained name crosswalk (製品名称読みかえ,
see name_crosswalk.py) rather than a regex/pattern heuristic — deliberately,
since a wrong stock match feeds a wrong order-quantity suggestion (real
money), while a missed match just means no suggestion appears at all. A
human-curated mapping is far less likely to produce the former. Verified
2026-07-07: 41/368 blank rows resolve to a real existing Item this way
(61,399 units of stock that would otherwise be invisible to the forecast).
"""
from __future__ import annotations

import glob
from datetime import date
from pathlib import Path

from ..config import SAP_STOCK_GLOB
from .excel import read_xlsx_rows
from .name_crosswalk import latest_crosswalk_file, load_crosswalk_tokens, resolve_via_crosswalk


def latest_sap_file() -> Path | None:
    hits = sorted(glob.glob(SAP_STOCK_GLOB))
    return Path(hits[-1]) if hits else None


def parse_sap_stock() -> dict:
    path = latest_sap_file()
    if path is None:
        return {"path": None, "rows": {}, "unnamed": 0, "resolved_via_crosswalk": 0}
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    crosswalk_path = latest_crosswalk_file()
    tokens = load_crosswalk_tokens(crosswalk_path) if crosswalk_path else []

    hdr = None
    qty = {}
    unnamed = 0
    resolved_via_crosswalk = 0
    for row in ws.iter_rows(values_only=True):
        if hdr is None:
            if row and any(v == "あだ名" for v in row):
                hdr = {str(v): i for i, v in enumerate(row) if v is not None}
            continue
        adana = row[hdr["あだ名"]] if "あだ名" in hdr else None
        syoname = row[hdr["SYONAME"]] if "SYONAME" in hdr else None
        q = row[hdr.get("QTYの合計", len(row) - 1)]
        if adana in (None, ""):
            match = resolve_via_crosswalk(syoname, tokens) if tokens else None
            if match and isinstance(q, (int, float)):
                qty[match] = qty.get(match, 0.0) + float(q)
                resolved_via_crosswalk += 1
                continue
            unnamed += 1
            continue
        if isinstance(q, (int, float)):
            qty[str(adana).strip()] = qty.get(str(adana).strip(), 0.0) + float(q)
    wb.close()
    return {"path": path, "rows": qty, "unnamed": unnamed,
            "resolved_via_crosswalk": resolved_via_crosswalk,
            "month": date.today().replace(day=1)}
