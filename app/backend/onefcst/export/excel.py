"""Excel export — "everything checkable in Excel" (POC_BUILD_PLAN §7).

Every dashboard view has a corresponding export here. These are always NEW
workbooks written to app/data/exports/ — never a source file. Staff who
don't trust a number on screen can open the same data in the tool they
already trust and verify it directly.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from ..config import EXPORT_DIR
from ..db import get_session
from ..models import Item, CompatRule, Adjustment, AuditLog
from ..verify.mismatch import compute_mismatch
from ..verify.backtest import run as run_backtest
from ..advisory.adjustments import adjustment_rate

HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = HEADER_FONT
        c.fill = HEADER_FILL


def _autosize(ws, widths):
    for col, w in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = w


def _new_wb(title: str):
    wb = openpyxl.Workbook()
    wb.active.title = title
    return wb, wb.active


def export_mismatch() -> Path:
    session = get_session()
    report = compute_mismatch(session)
    session.close()
    wb, ws = _new_wb("ミスマッチ検知")
    ws.append([f"部材ミスマッチ検知レポート — 対象月: {report['month']}"])
    ws.append([f"出力日時: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    _header(ws, ["ソケット群", "CPU需要", "MB需要", "差分", "差分率(%)",
                "専用部材", "引き取り責任", "重大度"])
    for g in report["groups"]:
        ws.append([g["socket"], g["cpu_demand"], g["mb_demand"], g["gap"], g["pct"],
                  "あり" if g["exclusive"] else "", "あり" if g["takeorpay_n1"] else "",
                  g["severity"]])
    _autosize(ws, (14, 12, 12, 12, 12, 10, 12, 10))
    out = EXPORT_DIR / f"ミスマッチ検知_{date.today():%Y%m%d}.xlsx"
    wb.save(out)
    return out


def export_compat_rules(status: str | None = None) -> Path:
    session = get_session()
    q = session.query(CompatRule)
    if status:
        q = q.filter(CompatRule.status == status)
    rows = q.order_by(CompatRule.confidence.desc().nulls_last()).all()
    items = {i.id: i.canonical_name for i in session.query(Item).all()}
    session.close()
    wb, ws = _new_wb("互換性ルール")
    ws.append([f"部材互換性ルール一覧（{status or '全件'}） — 出力日時: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    # 部材A/部材B rather than CPU/マザーボード: the schema is a generic
    # part_a/part_b pair, and every rule happens to be a CPU↔MB pair only
    # because that is all the drafting has covered so far. Hardcoded category
    # names would silently mislabel the first rule between any other pair.
    _header(ws, ["部材A", "部材B", "関係", "信頼度", "提案元", "状態", "承認者"])
    for r in rows:
        ws.append([items.get(r.part_a_id, "?"), items.get(r.part_b_id, "?"), r.relation,
                  r.confidence, r.origin, r.status, r.confirmed_by or ""])
    _autosize(ws, (18, 18, 10, 10, 16, 12, 16))
    # The status belongs in the FILENAME, not only in the title row: the UI
    # exports whatever filter is on screen, so without it a buyer who saves
    # drafts and then confirmed ends up with "…(1).xlsx" and no way to tell
    # the two apart outside the portal.
    label = {"draft": "未確認", "confirmed": "確認済", "rejected": "却下"}.get(
        status or "", "全件")
    out = EXPORT_DIR / f"互換性ルール_{label}_{date.today():%Y%m%d}.xlsx"
    wb.save(out)
    return out


def export_adjustments() -> Path:
    session = get_session()
    rows = session.query(Adjustment).order_by(Adjustment.created_at.desc()).all()
    items = {i.id: i.canonical_name for i in session.query(Item).all()}
    rate = adjustment_rate(session)
    session.close()
    wb, ws = _new_wb("調整履歴")
    ws.append([f"予測調整履歴 — 出力日時: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([f"調整率: {rate['adjusted_cells']}/{rate['total_forecast_cells']} "
              f"({rate['rate_pct']}%)" if rate["rate_pct"] is not None else "調整率: —"])
    ws.append([])
    _header(ws, ["対象部材", "対象月", "範囲", "変更前", "変更後", "理由", "担当者", "記録日時"])
    for a in rows:
        ws.append([items.get(a.item_id, "?"), a.month.isoformat() if a.month else "",
                  a.scope, a.old_value, a.new_value, a.reason, a.author,
                  a.created_at.strftime("%Y-%m-%d %H:%M")])
    _autosize(ws, (18, 12, 10, 12, 12, 40, 16, 18))
    out = EXPORT_DIR / f"調整履歴_{date.today():%Y%m%d}.xlsx"
    wb.save(out)
    return out


def export_backtest() -> Path:
    report = run_backtest()
    wb, ws = _new_wb("予測精度検証")
    ws.append([f"予測精度バックテスト結果 — 出力日時: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append(["機械学習は使用せず、統計的季節性モデルを実データで検証した結果です。"])
    ws.append([])
    _header(ws, ["期", "実績合計", "季節性モデル WAPE(%)", "均等配分 WAPE(%)", "前年同月 WAPE(%)"])
    for r in report["rows"]:
        ws.append([f"{r['ki']}期", r["total"], r["seasonal_wape"], r["flat_wape"],
                  r["last_year_wape"]])
    ws.append([])
    ws.append(["平均", "", report["avg_seasonal_wape"], report["avg_flat_wape"],
              report["avg_last_year_wape"]])
    ws.append([f"最良手法（アウトオブサンプル）: {report['winner']}"])
    _autosize(ws, (10, 14, 18, 16, 16))
    out = EXPORT_DIR / f"予測精度検証_{date.today():%Y%m%d}.xlsx"
    wb.save(out)
    return out


def export_audit_month(year: int, month: int, out_dir: Path | None = None) -> tuple[Path, int]:
    """One calendar month of the audit log, as the monthly report.

    Returns (path, row_count) — the count is what `audit_archive` verifies the
    written file against before it is allowed to delete anything.

    Differs from export_audit_log() below in the two ways that matter for a
    report someone files: it is bounded by MONTH rather than "latest 500", so
    the file is a complete record of a period rather than a snapshot; and it
    resolves each author's name and department from user_profiles, because a
    report read a year later by someone in another department needs to say
    「山田 太郎 / 購買本部」, not `m_yamada@twave.co.jp`.
    """
    from ..models import UserProfile

    start = datetime(year, month, 1)
    end = datetime(year + (month == 12), (month % 12) + 1, 1)

    session = get_session()
    try:
        rows = (session.query(AuditLog)
                .filter(AuditLog.at >= start, AuditLog.at < end)
                .order_by(AuditLog.at.asc()).all())
        # Ascending: a report is read start-of-month first, unlike the screen
        # feed which shows newest first.
        people = {p.email: p for p in session.query(UserProfile).all()}
    finally:
        session.close()

    wb, ws = _new_wb("変更履歴")
    ws.append([f"操作・変更履歴 {year}年{month}月 — 出力日時: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([f"対象期間: {start:%Y-%m-%d} 〜 {end - timedelta(days=1):%Y-%m-%d}　件数: {len(rows)}"])
    ws.append([])
    _header(ws, ["日時", "担当者", "部署", "メールアドレス", "内容", "変更前", "変更後"])
    for r in rows:
        p = people.get((r.who or "").lower())
        ws.append([r.at.strftime("%Y-%m-%d %H:%M:%S"),
                  (p.display_name if p and p.display_name else (r.who or "system")),
                  (p.department if p and p.department else ""),
                  r.who or "",
                  r.what, (r.before or "")[:200], (r.after or "")[:200]])
    _autosize(ws, (20, 18, 18, 28, 44, 30, 30))

    out_dir = out_dir or (EXPORT_DIR / "audit_archive")
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"変更履歴_{year}{month:02d}.xlsx"
    wb.save(out)
    return out, len(rows)


def export_audit_log(limit: int = 500) -> Path:
    session = get_session()
    rows = session.query(AuditLog).order_by(AuditLog.at.desc()).limit(limit).all()
    session.close()
    wb, ws = _new_wb("変更履歴")
    ws.append([f"操作・変更履歴（直近{limit}件） — 出力日時: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    _header(ws, ["日時", "担当者", "内容", "変更前", "変更後"])
    for r in rows:
        ws.append([r.at.strftime("%Y-%m-%d %H:%M:%S"), r.who or "system", r.what,
                  (r.before or "")[:200], (r.after or "")[:200]])
    _autosize(ws, (18, 16, 40, 40, 40))
    out = EXPORT_DIR / f"変更履歴_{date.today():%Y%m%d}.xlsx"
    wb.save(out)
    return out


# Ownership-tier labels for the forecast export (Item.tier: 1 = the four
# categories 企画部 plans directly, 2 = everything purchasing derives).
# Keep in sync with config.TIER1_CATEGORIES.
_TIER_SCOPE = {
    None: ("全部材", "全部材（企画4部材＋購買部材）"),
    1: ("企画4部材", "企画部門の4部材（CPU・グラフィックカード・マザーボード・PCケース）"),
    2: ("購買部材", "購買部門の部材（企画4部材以外すべて）"),
}


def export_forecasts(tier: int | None = None) -> Path:
    """需要予測一覧 — every SKU's monthly baseline for the latest run, plus
    stock, method, hold-out accuracy, and the order suggestion if any.

    tier: None = all items, 1 = planning-owned CPU/VGA/MB/Case only,
    2 = purchasing-owned everything else (same split as Item.tier)."""
    from .. import forecast_views
    session = get_session()
    data = forecast_views.list_forecasts(session, tier=tier, limit=10000)
    session.close()
    short, scope_label = _TIER_SCOPE[tier]
    wb, ws = _new_wb("需要予測")
    run = data.get("run") or {}
    ws.append([f"需要予測（部品別・月別数量） — 対象: {scope_label} — "
               f"実行: {run.get('created_at', '—')} "
               f"/ 手法: {run.get('method_version', '—')}"])
    ws.append([f"出力日時: {datetime.now():%Y-%m-%d %H:%M} — "
               "検証WAPEは各SKUの直近6か月ホールドアウトによる（低いほど良い）"])
    ws.append([])
    month_cols = [m[:7] for m in data["months"]]
    _header(ws, ["部材（綽名）", "カテゴリ", "手法", "検証WAPE(%)", "現在庫",
                 "6か月合計", "発注提案数量", "発注期限"] + month_cols)
    for r in data["rows"]:
        order = r.get("order") or {}
        ws.append([r["name"], r["category"] or "", r["method"] or "",
                   r["wape"], r["stock"], r["next6_total"],
                   order.get("qty"), (order.get("order_by") or "")[:7]]
                  + [q for q in r["qty"]])
    _autosize(ws, (22, 14, 10, 12, 10, 12, 12, 12))
    out = EXPORT_DIR / f"需要予測_{short}_{date.today():%Y%m%d}.xlsx"
    wb.save(out)
    return out
