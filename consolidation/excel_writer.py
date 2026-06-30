"""
Excel-export writer (Face 2 of the consolidation).

Reads consolidation/consolidated.db and produces a single formatted .xlsx that
rolls the 17 ONE files into 内田さん's master shape. This is the "download as
Excel" deliverable AND the in-place automation output — same code path.

Three sheets:
  1. サマリー        — per-category roll-up (SKUs / excess / low / negative)
  2. 統合一覧        — every SKU across all 17 files, one row each, with flags
  3. 照合レポート    — the reconciliation/accuracy report, so the file carries
                       its own proof that it faithfully matches source

If the DB has no data yet, it runs the consolidation first so export is one-click.

Run:  python3 excel_writer.py            → writes 統合出力_YYYYMMDD.xlsx
      python3 excel_writer.py out.xlsx   → custom path
"""
import sqlite3
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import consolidate

DB_PATH = consolidate.DB_PATH
HERE = Path(__file__).parent

# ── styling palette ───────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")   # dark blue
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT  = Font(bold=True, size=14, color="1F4E78")
SUB_FONT    = Font(italic=True, size=9, color="808080")

FILL_EXCESS = PatternFill("solid", fgColor="FCE4E4")   # light red
FILL_LOW    = PatternFill("solid", fgColor="FFF2CC")   # light amber
FILL_NEG    = PatternFill("solid", fgColor="F4B6B6")   # stronger red
FILL_PASS   = PatternFill("solid", fgColor="E2EFDA")   # light green
FILL_WARN   = PatternFill("solid", fgColor="FFF2CC")
FILL_FAIL   = PatternFill("solid", fgColor="F4B6B6")

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center")


def _ensure_data(conn):
    """If the consolidated table is empty/missing, build it first."""
    try:
        n = conn.execute("SELECT COUNT(*) FROM consolidated_skus").fetchone()[0]
    except sqlite3.OperationalError:
        n = 0
    if n == 0:
        print("  consolidated.db empty — running consolidation first...")
        consolidate.run()


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Sheet 1: summary roll-up ──────────────────────────────────────────────────

def build_summary(wb, conn, snapshot):
    ws = wb.active
    ws.title = "サマリー"
    ws["A1"] = "ONE統合サマリー（全17ファイル）"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"スナップショット日: {snapshot}　／　自動生成: 統合エンジン"
    ws["A2"].font = SUB_FONT

    headers = ["カテゴリ", "ファイル", "SKU数", "🔴 過剰", "⚠️ 不足", "❗ マイナス在庫"]
    hrow = 4
    for j, h in enumerate(headers, start=1):
        ws.cell(row=hrow, column=j, value=h)
    _style_header(ws, hrow, len(headers))

    rows = conn.execute("""
        SELECT category, source_file,
               COUNT(*), SUM(flag_excess), SUM(flag_low), SUM(flag_negative)
        FROM consolidated_skus GROUP BY category, source_file
        ORDER BY COUNT(*) DESC
    """).fetchall()

    r = hrow + 1
    for cat, src, n, exc, low, neg in rows:
        vals = [cat, src, n, exc or 0, low or 0, neg or 0]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = BORDER
            cell.alignment = LEFT if j <= 2 else CENTER
        if (exc or 0) > 0:
            ws.cell(row=r, column=4).fill = FILL_EXCESS
        if (neg or 0) > 0:
            ws.cell(row=r, column=6).fill = FILL_NEG
        r += 1

    # totals row
    tot = conn.execute("""SELECT COUNT(*), SUM(flag_excess), SUM(flag_low),
                          SUM(flag_negative) FROM consolidated_skus""").fetchone()
    ws.cell(row=r, column=1, value="合計")
    ws.cell(row=r, column=3, value=tot[0])
    ws.cell(row=r, column=4, value=tot[1] or 0)
    ws.cell(row=r, column=5, value=tot[2] or 0)
    ws.cell(row=r, column=6, value=tot[3] or 0)
    for j in range(1, len(headers) + 1):
        c = ws.cell(row=r, column=j)
        c.font = Font(bold=True)
        c.border = BORDER
        c.alignment = LEFT if j <= 2 else CENTER

    _autofit(ws, [12, 30, 8, 9, 9, 14])
    ws.freeze_panes = "A5"


# ── Sheet 2: full consolidated list ───────────────────────────────────────────

def build_detail(wb, conn):
    ws = wb.create_sheet("統合一覧")
    headers = ["カテゴリ", "SKU（綽名）", "ソースファイル", "安全在庫(日)",
               "在庫残高", "在庫日数", "状態", "スナップショット日"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h)
    _style_header(ws, 1, len(headers))

    rows = conn.execute("""
        SELECT category, sku, source_file, safety_days, stock_balance,
               stock_days, flag_excess, flag_low, flag_negative, snapshot_date
        FROM consolidated_skus
        ORDER BY category, stock_days DESC
    """).fetchall()

    r = 2
    for (cat, sku, src, safety, bal, days, exc, low, neg, snap) in rows:
        status = []
        if neg: status.append("マイナス")
        if exc: status.append("過剰")
        if low: status.append("不足")
        status_txt = "・".join(status) if status else "OK"

        vals = [cat, sku, src,
                round(safety, 1) if safety is not None else None,
                round(bal, 1) if bal is not None else None,
                round(days, 1) if days is not None else None,
                status_txt, snap]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = BORDER
            cell.alignment = LEFT if j in (1, 2, 3, 7, 8) else CENTER

        # row highlight by severity (negative > excess > low)
        fill = FILL_NEG if neg else FILL_EXCESS if exc else FILL_LOW if low else None
        if fill:
            for j in range(1, len(headers) + 1):
                ws.cell(row=r, column=j).fill = fill
        r += 1

    _autofit(ws, [11, 34, 26, 11, 11, 10, 14, 16])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{r-1}"


# ── Sheet 3: reconciliation / accuracy report ─────────────────────────────────

def build_reconciliation(wb, conn):
    ws = wb.create_sheet("照合レポート")
    ws["A1"] = "照合レポート（正確性ゲート）"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "各ファイルの抽出結果を期待SKU数と突合。PASS=完全一致。"
    ws["A2"].font = SUB_FONT

    try:
        recon = conn.execute("""
            SELECT category, source_file, skus_parsed, expected_skus,
                   blocks_seen, status, note FROM reconciliation
            ORDER BY status DESC, category
        """).fetchall()
    except sqlite3.OperationalError:
        recon = []

    if not recon:
        ws["A4"] = "照合データなし。reconcile.py を実行してください。"
        return

    headers = ["カテゴリ", "ファイル", "抽出SKU", "期待SKU", "ブロック数", "状態", "備考"]
    hrow = 4
    for j, h in enumerate(headers, start=1):
        ws.cell(row=hrow, column=j, value=h)
    _style_header(ws, hrow, len(headers))

    r = hrow + 1
    n_pass = n_warn = n_fail = 0
    for (cat, src, parsed, expected, blocks, status, note) in recon:
        n_pass += status == "PASS"; n_warn += status == "WARN"; n_fail += status == "FAIL"
        vals = [cat, src, parsed, expected, blocks, status, note or ""]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = BORDER
            cell.alignment = LEFT if j in (1, 2, 7) else CENTER
        sf = {"PASS": FILL_PASS, "WARN": FILL_WARN, "FAIL": FILL_FAIL}.get(status)
        if sf:
            ws.cell(row=r, column=6).fill = sf
        r += 1

    total_parsed = sum(x[2] for x in recon)
    total_exp = sum(x[3] for x in recon)
    ws.cell(row=r + 1, column=1, value="判定")
    ws.cell(row=r + 1, column=1).font = Font(bold=True)
    verdict = (f"✓ 合格 — {total_parsed}/{total_exp} SKU 抽出 (100%)　"
               f"{n_pass} PASS / {n_warn} WARN / {n_fail} FAIL"
               if n_fail == 0 and total_parsed == total_exp
               else f"⚠ 要確認 — {total_parsed}/{total_exp} SKU　"
                    f"{n_pass} PASS / {n_warn} WARN / {n_fail} FAIL")
    ws.cell(row=r + 1, column=2, value=verdict)
    ws.cell(row=r + 1, column=2).font = Font(bold=True,
            color="375623" if n_fail == 0 else "9C5700")

    _autofit(ws, [12, 30, 9, 9, 11, 8, 40])
    ws.freeze_panes = "A5"


# ── main ──────────────────────────────────────────────────────────────────────

def export(out_path=None):
    conn = sqlite3.connect(DB_PATH)
    _ensure_data(conn)

    snapshot = conn.execute(
        "SELECT MAX(snapshot_date) FROM consolidated_skus").fetchone()[0] or date.today().isoformat()

    if out_path is None:
        stamp = date.today().strftime("%Y%m%d")
        out_path = HERE / f"統合出力_{stamp}.xlsx"
    out_path = Path(out_path)

    wb = openpyxl.Workbook()
    build_summary(wb, conn, snapshot)
    build_detail(wb, conn)
    build_reconciliation(wb, conn)
    wb.save(out_path)
    conn.close()

    print(f"✓ Excel書き出し完了 → {out_path}")
    print(f"  シート: サマリー / 統合一覧 / 照合レポート")
    return out_path


if __name__ == "__main__":
    export(sys.argv[1] if len(sys.argv) > 1 else None)
