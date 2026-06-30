"""
Consolidate 部品詳細 from all 17 product ONE files into ONE unified SKU table.

Parameterized by per-file label_col / name_col / first_date_col (config_all.py),
so the two non-standard layouts (FAN, Case_梱包) parse correctly alongside the
15 standard files. Writes a single consolidated table to consolidation/consolidated.db
and returns per-file parse stats for the reconciliation check (reconcile.py).

This is deliberately SEPARATE from poc/extractor.py so the 5-anchor forecast
pipeline is never touched.
"""
import sqlite3
from datetime import datetime, timedelta, date
from pathlib import Path
import pyxlsb
import openpyxl
from config_all import FILES, BLOCK_LABELS, EXCESS_DAYS, BASE

DB_PATH = Path(__file__).parent / "consolidated.db"


# ── value helpers ─────────────────────────────────────────────────────────────

def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        f = float(v)
        if f <= 0:
            return None
        d = (datetime(1899, 12, 30) + timedelta(days=f)).date()
        return d if 2010 <= d.year <= 2035 else None
    except (TypeError, ValueError):
        return None


def to_float(v):
    if v is None:
        return None
    if isinstance(v, str) and (v.startswith("#") or v.strip() == ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ── sheet reader (both formats) ───────────────────────────────────────────────

def read_sheet(path, fmt, sheet="部品詳細"):
    """Return (rows, sheet_present, all_sheet_names). rows = list[{col: val}]."""
    rows = []
    if fmt == "xlsb":
        with pyxlsb.open_workbook(str(path)) as wb:
            names = wb.sheets
            if sheet not in names:
                return [], False, names
            with wb.get_sheet(sheet) as ws:
                for row in ws.rows():
                    rows.append({c.c: c.v for c in row if c.v is not None})
        return rows, True, names
    else:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        names = wb.sheetnames
        if sheet not in names:
            wb.close()
            return [], False, names
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            rows.append({j: v for j, v in enumerate(row) if v is not None})
        wb.close()
        return rows, True, names


# ── 部品詳細 parser (parameterized) ───────────────────────────────────────────

def parse_buhin_detail(cat, cfg, snapshot_date):
    """
    Parse 部品詳細 into per-SKU records using this file's column layout.
    Returns (records, stats) where stats feeds the reconciliation check.
    """
    path = Path(cfg.get("_abspath") or (BASE / cfg["file"]))
    label_col = cfg["label_col"]
    name_col  = cfg["name_col"]
    fdc       = cfg["first_date_col"]

    stats = {"category": cat, "file": cfg["file"], "sheet_found": False,
             "blocks_seen": 0, "skus_parsed": 0, "skipped_no_name": 0,
             "snapshot_col": None, "snapshot_date": None, "snapshot_stale": False,
             "note": ""}

    try:
        rows, present, names = read_sheet(path, cfg["fmt"])
    except Exception as e:
        stats["note"] = f"READ ERROR: {e}"
        return [], stats
    if not present:
        stats["note"] = f"NO 部品詳細 sheet (sheets: {names[:6]})"
        return [], stats
    if not rows:
        stats["note"] = "EMPTY sheet"
        return [], stats
    stats["sheet_found"] = True

    # date header (row 0), from first_date_col onward
    header = rows[0]
    date_cols = {}
    for c, v in header.items():
        if c < fdc:
            continue
        d = to_date(v)
        if d:
            date_cols[c] = d
    if not date_cols:
        stats["note"] = "no date columns in header row 0"
        return [], stats

    sorted_dates = sorted(date_cols.items())  # [(col, date)]
    # snapshot col = first date >= snapshot_date, else the last available
    snap_col = next((c for c, d in sorted_dates if d >= snapshot_date), None)
    if snap_col is None:
        # No column reaches snapshot_date: file is stale. Fall back to the last
        # column, but FLAG it — that value is the last planned balance, not a
        # current one, so downstream must not treat it as "today's stock".
        snap_col = sorted_dates[-1][0]
        stats["snapshot_stale"] = True
        stats["note"] = (f"STALE: no date >= {snapshot_date}; using last column "
                         f"{date_cols[snap_col].isoformat()} (planned, not current)")
    else:
        stats["snapshot_stale"] = False
    stats["snapshot_col"] = snap_col
    stats["snapshot_date"] = date_cols[snap_col].isoformat()

    records = []
    i = 0
    while i + 3 < len(rows):
        r0, r1, r2, r3 = rows[i], rows[i + 1], rows[i + 2], rows[i + 3]
        if str(r0.get(label_col, "")).strip() != BLOCK_LABELS[0]:
            i += 1
            continue
        # validate the 4-row block
        if (str(r1.get(label_col, "")).strip() != BLOCK_LABELS[1] or
                str(r2.get(label_col, "")).strip() != BLOCK_LABELS[2] or
                str(r3.get(label_col, "")).strip() != BLOCK_LABELS[3]):
            i += 1
            continue

        stats["blocks_seen"] += 1
        sku = str(r0.get(name_col, "")).strip()
        if not sku:
            stats["skipped_no_name"] += 1
            i += 4
            continue

        safety_days   = to_float(r1.get(name_col))      # safety target in 購買計画 row, name_col
        stock_balance = to_float(r2.get(snap_col))
        stock_days    = to_float(r3.get(snap_col))

        excess_thr = EXCESS_DAYS.get(cat, EXCESS_DAYS["_default"])
        is_excess  = bool(stock_days is not None and stock_days > excess_thr)
        is_low     = bool(stock_days is not None and safety_days and stock_days < safety_days * 0.5)
        neg        = bool(stock_balance is not None and stock_balance < 0)

        records.append({
            "category":       cat,
            "source_file":    cfg["file"],
            "sku":            sku,
            "snapshot_date":  date_cols[snap_col].isoformat(),
            "safety_days":    safety_days,
            "stock_balance":  stock_balance,
            "stock_days":     stock_days,
            "flag_excess":    int(is_excess),
            "flag_low":       int(is_low),
            "flag_negative":  int(neg),
        })
        stats["skus_parsed"] += 1
        i += 4

    return records, stats


# ── DB ────────────────────────────────────────────────────────────────────────

def init_db(conn):
    conn.executescript("""
    DROP TABLE IF EXISTS consolidated_skus;
    CREATE TABLE consolidated_skus (
        id INTEGER PRIMARY KEY,
        category TEXT, source_file TEXT, sku TEXT, snapshot_date TEXT,
        safety_days REAL, stock_balance REAL, stock_days REAL,
        flag_excess INTEGER, flag_low INTEGER, flag_negative INTEGER
    );
    DROP TABLE IF EXISTS reconciliation;
    CREATE TABLE reconciliation (
        id INTEGER PRIMARY KEY,
        category TEXT, source_file TEXT, sheet_found INTEGER,
        blocks_seen INTEGER, skus_parsed INTEGER, expected_skus INTEGER,
        skipped_no_name INTEGER, snapshot_date TEXT,
        status TEXT, note TEXT
    );
    """)
    conn.commit()


def run(snapshot_date=None, db_path=DB_PATH):
    if snapshot_date is None:
        snapshot_date = date.today()

    conn = sqlite3.connect(db_path)
    init_db(conn)

    all_stats = []
    for cat, cfg in FILES.items():
        recs, stats = parse_buhin_detail(cat, cfg, snapshot_date)
        stats["expected_skus"] = cfg["expected_skus"]
        if recs:
            cols = list(recs[0].keys())
            ph = ",".join("?" * len(cols))
            conn.executemany(
                f"INSERT INTO consolidated_skus ({','.join(cols)}) VALUES ({ph})",
                [tuple(r[c] for c in cols) for r in recs],
            )
        all_stats.append(stats)
        print(f"  [{cat:<11}] parsed {stats['skus_parsed']:>3} / {cfg['expected_skus']:>3} expected"
              f"  {stats['note']}")
    conn.commit()
    conn.close()
    return all_stats


if __name__ == "__main__":
    print(f"Consolidating 17 files (snapshot {date.today()})...")
    run()
    print(f"\nWrote → {DB_PATH}")
    print("Run `python3 reconcile.py` for the accuracy report.")
