"""
購買計画-based Phase-1 features (NAS-Excel only, no forecasting):
  • Overdue / delayed PO detection      (survey 4.0)
  • Price / supplier history lookup     (survey 4.3)
  • Incoming-stock → excess warning     (survey 4.5)

Parses 購買計画 across all 17 files by AUTO-DETECTING columns from the header
row (数量 / $価格 / ¥価格 / 納品日 / IO / 発注日 / 綽名 / 仕入先 / 発注番号),
so it is robust to per-file column shifts. Status (納品済/キャンセル) is read
from the row text. Writes result tables into consolidated.db.

Run:  python3 purchasing.py
"""
import os
import sqlite3
from datetime import datetime, timedelta, date

import pyxlsb
import openpyxl

import consolidate
from config_all import FILES, BASE, EXCESS_DAYS

DB_PATH = consolidate.DB_PATH
# "Today" for overdue/incoming windows. Defaults to the real current date so the
# 180-day overdue window and future-inbound checks stay correct over time;
# override via PROCUREMENT_TODAY=YYYY-MM-DD for reproducible/test runs.
TODAY = (date.fromisoformat(os.environ["PROCUREMENT_TODAY"])
         if os.environ.get("PROCUREMENT_TODAY") else date.today())


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        f = float(v)
        if f <= 0:
            return None
        d = (datetime(1899, 12, 30) + timedelta(days=f)).date()
        return d if 2010 <= d.year <= 2035 else None
    except (TypeError, ValueError):
        return None


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def cell_str(v):
    """Trimmed string, or None for empty/None (avoids the str(None)=='None' trap)."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _detect_cols(header):
    """Map field -> column index from the 購買計画 header row."""
    col = {}
    for c, raw in header.items():
        if raw is None:
            continue
        h = str(raw).strip()
        if "数量" in h:                                   col["qty"] = c
        elif "$" in h and ("価格" in h or "単価" in h):    col["usd"] = c
        elif any(y in h for y in ("¥", "\\", "￥")) and ("価格" in h or "単価" in h): col["jpy"] = c
        elif "納品" in h and "済" not in h:                col["delivery"] = c
        elif h == "IO":                                   col["io"] = c
        elif "発注日" in h:                                col["order_date"] = c
        elif "綽名" in h:                                  col["sku"] = c
        elif ("仕入先" in h or "購買先" in h) and "コード" not in h: col["supplier"] = c
        elif "発注書番号" in h or "発注番号" in h:          col["po_num"] = c
    return col


def read_kobai(cat):
    """Return list of PO dicts for one file's 購買計画 (auto-detected layout)."""
    cfg = FILES[cat]
    rows = []
    if cfg["fmt"] == "xlsb":
        with pyxlsb.open_workbook(str(BASE / cfg["file"])) as wb:
            if "購買計画" not in wb.sheets:
                return []
            with wb.get_sheet("購買計画") as ws:
                for r in ws.rows():
                    rows.append({c.c: c.v for c in r})
    else:
        wb = openpyxl.load_workbook(str(BASE / cfg["file"]), data_only=True)
        if "購買計画" not in wb.sheetnames:
            return []
        ws = wb["購買計画"]
        for r in ws.iter_rows(values_only=False):
            rows.append({c.column - 1: c.value for c in r})
    if not rows:
        return []

    col = _detect_cols(rows[0])
    if "sku" not in col or "delivery" not in col:
        return []   # can't use this file's PO sheet reliably

    out = []
    for row in rows[1:]:
        sku = cell_str(row.get(col["sku"]))
        if not sku:
            continue
        rowtext = " ".join(str(v) for v in row.values() if isinstance(v, str))
        delivered = "納品済" in rowtext
        cancelled = "キャンセル" in rowtext
        out.append({
            "category": cat, "file": cfg["file"], "sku": sku,
            "qty": to_float(row.get(col.get("qty"))),
            "usd": to_float(row.get(col.get("usd"))),
            "jpy": to_float(row.get(col.get("jpy"))),
            "delivery": to_date(row.get(col.get("delivery"))),
            "order_date": to_date(row.get(col.get("order_date"))),
            "io": to_float(row.get(col.get("io"))),
            "supplier": cell_str(row.get(col.get("supplier"))),
            "po_num": cell_str(row.get(col.get("po_num"))),
            "status": "delivered" if delivered else "cancelled" if cancelled else "open",
        })
    return out


def read_all_pos():
    pos = []
    for cat in FILES:
        try:
            pos += read_kobai(cat)
        except Exception:
            pass
    return pos


# ── feature 1: overdue / delayed POs ──────────────────────────────────────────

OVERDUE_WINDOW = 180   # only POs overdue within the last N days are actionable


def overdue(pos):
    """Actionable overdue: open PO whose delivery date passed within the last
    OVERDUE_WINDOW days. Older 'open' rows are stale records (un-annotated), not
    real delays, so they are excluded."""
    out = []
    lo = TODAY - timedelta(days=OVERDUE_WINDOW)
    for p in pos:
        # actionable = an actually-placed order (has PO# + supplier), open, recently overdue
        if (p["status"] == "open" and p["po_num"] and p["supplier"]
                and p["delivery"] and lo <= p["delivery"] < TODAY):
            days = (TODAY - p["delivery"]).days
            out.append({**p, "overdue_days": days})
    out.sort(key=lambda x: -x["overdue_days"])
    return out


# ── feature 2: price / supplier history ───────────────────────────────────────

def price_history(pos):
    """Per SKU: latest price/supplier + full history (rows with a price)."""
    by_sku = {}
    for p in pos:
        if p["usd"] is None and p["jpy"] is None:
            continue
        by_sku.setdefault((p["category"], p["sku"]), []).append(p)
    rows = []
    for (cat, sku), lst in by_sku.items():
        lst.sort(key=lambda x: x["order_date"] or date(1900, 1, 1))
        latest = lst[-1]
        rows.append({"category": cat, "sku": sku, "orders": len(lst),
                     "latest_usd": latest["usd"], "latest_jpy": latest["jpy"],
                     "latest_supplier": latest["supplier"],
                     "latest_date": latest["order_date"].isoformat() if latest["order_date"] else None,
                     "suppliers": ", ".join(sorted({p["supplier"] for p in lst if p["supplier"]}))})
    rows.sort(key=lambda x: (x["category"], x["sku"]))
    return rows


# ── feature 3: incoming-stock → excess warning ────────────────────────────────

def incoming_excess(pos):
    """Open future inbound that would push a SKU past its excess-days threshold."""
    conn = sqlite3.connect(DB_PATH)
    stock = {(c, s): (b, d) for c, s, b, d in conn.execute(
        "SELECT category, sku, stock_balance, stock_days FROM consolidated_skus")}
    conn.close()

    inbound = {}
    for p in pos:
        if p["status"] == "open" and p["qty"] and p["delivery"] and p["delivery"] >= TODAY:
            inbound[(p["category"], p["sku"])] = inbound.get((p["category"], p["sku"]), 0) + p["qty"]

    out = []
    for (cat, sku), qty in inbound.items():
        bal, days = stock.get((cat, sku), (None, None))
        if bal is None or days is None or bal <= 0:
            continue
        thr = EXCESS_DAYS.get(cat, EXCESS_DAYS["_default"])
        days_after = days * (bal + qty) / bal       # project days once inbound lands
        if days_after > thr:
            out.append({"category": cat, "sku": sku, "stock_days": round(days),
                        "inbound_qty": round(qty), "days_after": round(days_after),
                        "threshold": thr, "already_excess": days > thr})
    out.sort(key=lambda x: -x["days_after"])
    return out


def build_portal_data():
    """Read the three PO result tables for the portal (builds them if missing)."""
    conn = sqlite3.connect(DB_PATH)

    def q(sql):
        try:
            return conn.execute(sql).fetchall()
        except sqlite3.OperationalError:
            return None

    ov = q("SELECT category,sku,qty,delivery,overdue_days,supplier,po_num FROM po_overdue ORDER BY overdue_days DESC")
    if ov is None:
        conn.close(); run(); conn = sqlite3.connect(DB_PATH)
        ov = conn.execute("SELECT category,sku,qty,delivery,overdue_days,supplier,po_num FROM po_overdue ORDER BY overdue_days DESC").fetchall()
    ie = conn.execute("SELECT category,sku,stock_days,inbound_qty,days_after,threshold,already_excess FROM po_incoming ORDER BY days_after DESC").fetchall()
    ph = conn.execute("SELECT category,sku,orders,latest_usd,latest_jpy,latest_supplier,latest_date,suppliers FROM po_price ORDER BY category,sku").fetchall()
    conn.close()

    return {
        "counts": {"overdue": len(ov), "incoming": len(ie), "price": len(ph)},
        "overdue": [{"cat": r[0], "sku": r[1], "qty": r[2], "delivery": r[3],
                     "days": r[4], "supplier": r[5], "po": r[6]} for r in ov],
        "incoming": [{"cat": r[0], "sku": r[1], "stock_days": r[2], "inbound": r[3],
                      "days_after": r[4], "threshold": r[5], "already": bool(r[6])} for r in ie],
        "price": [{"cat": r[0], "sku": r[1], "orders": r[2], "usd": r[3], "jpy": r[4],
                   "supplier": r[5], "date": r[6], "suppliers": r[7]} for r in ph],
    }


def run():
    pos = read_all_pos()
    ov = overdue(pos)
    ph = price_history(pos)
    ie = incoming_excess(pos)

    conn = sqlite3.connect(DB_PATH)
    conn.executescript("""
    DROP TABLE IF EXISTS po_overdue; DROP TABLE IF EXISTS po_price; DROP TABLE IF EXISTS po_incoming;
    CREATE TABLE po_overdue (category TEXT, sku TEXT, qty REAL, delivery TEXT, order_date TEXT,
                             supplier TEXT, po_num TEXT, overdue_days INTEGER);
    CREATE TABLE po_price (category TEXT, sku TEXT, orders INTEGER, latest_usd REAL, latest_jpy REAL,
                           latest_supplier TEXT, latest_date TEXT, suppliers TEXT);
    CREATE TABLE po_incoming (category TEXT, sku TEXT, stock_days INTEGER, inbound_qty INTEGER,
                              days_after INTEGER, threshold INTEGER, already_excess INTEGER);
    """)
    conn.executemany("INSERT INTO po_overdue VALUES (?,?,?,?,?,?,?,?)",
                     [(p["category"], p["sku"], p["qty"],
                       p["delivery"].isoformat() if p["delivery"] else None,
                       p["order_date"].isoformat() if p["order_date"] else None,
                       p["supplier"], p["po_num"], p["overdue_days"]) for p in ov])
    conn.executemany("INSERT INTO po_price VALUES (?,?,?,?,?,?,?,?)",
                     [(r["category"], r["sku"], r["orders"], r["latest_usd"], r["latest_jpy"],
                       r["latest_supplier"], r["latest_date"], r["suppliers"]) for r in ph])
    conn.executemany("INSERT INTO po_incoming VALUES (?,?,?,?,?,?,?)",
                     [(r["category"], r["sku"], r["stock_days"], r["inbound_qty"],
                       r["days_after"], r["threshold"], 1 if r["already_excess"] else 0) for r in ie])
    conn.commit(); conn.close()

    print(f"=== 購買計画 features — {TODAY} ===")
    print(f"  POs parsed: {len(pos)}")
    print(f"  ⏰ overdue POs: {len(ov)}")
    print(f"  💱 SKUs with price history: {len(ph)}")
    print(f"  📦 incoming→excess warnings: {len(ie)}")
    return ov, ph, ie


if __name__ == "__main__":
    run()
